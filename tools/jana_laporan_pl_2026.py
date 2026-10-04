"""Jana 'Profit & Loss Report 2026' (sistem input harian) dalam Excel.

Guna:
    python tools/jana_laporan_pl_2026.py OUTPUT.xlsx [FAIL_LAMA.xlsx]

Jika FAIL_LAMA (format manual: sheet 'JAN 2026', 'FEB 2026', ...) diberi,
data input harian di dalamnya dipindahkan ke sheet INPUT HARIAN.
"""
import calendar
import datetime as dt
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

TAHUN = 2026
N_INPUT = int(__import__("os").environ.get("N_INPUT", 2500))  # baris input harian
N_RETURN = 400
N_PROD = 4          # slot produk
N_CH = 5            # slot saluran
N_RATE = 40         # slot kadar harga
BULAN = ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY", "AUGUST",
         "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER"]
PAKEJ = ["TRIAL", "JIMAT", "SEMBUH"]

# ---------------- gaya ----------------
F = "Arial"
NAVY, TEAL, GREEN = "1F3A5F", "2A7F8F", "3A7D44"
f_title = Font(name=F, size=16, bold=True, color=NAVY)
f_sub = Font(name=F, size=9, italic=True, color="666666")
f_head = Font(name=F, size=9, bold=True, color="FFFFFF")
f_body = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_in = Font(name=F, size=10, color="0000FF")
f_sec = Font(name=F, size=12, bold=True, color=TEAL)
fill_in_h = PatternFill("solid", fgColor=NAVY)
fill_auto_h = PatternFill("solid", fgColor=TEAL)
fill_in = PatternFill("solid", fgColor="FFF6CC")
fill_auto = PatternFill("solid", fgColor="EEF3F8")
fill_tot = PatternFill("solid", fgColor="D9E2EC")
thin = Side(style="thin", color="C8D1DC")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
RM = '#,##0.00;[Red]-#,##0.00;"-"'
NUM = '#,##0;[Red]-#,##0;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
X2 = '0.00;[Red]-0.00;"-"'
DATE = "dd/mm/yyyy"
red = Font(name=F, size=10, bold=True, color="C00000")


def cell(ws, r, c, v=None, fmt=None, font=f_body, fill=None, align=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font, x.border = font, box
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    if align:
        x.alignment = align
    return x


def head(ws, r, c, v, fill=fill_in_h):
    return cell(ws, r, c, v, font=f_head, fill=fill, align=center)


def title(ws, t, s):
    ws["A1"], ws["A2"] = t, s
    ws["A1"].font, ws["A2"].font = f_title, f_sub
    ws.sheet_view.showGridLines = False


def dv_list(ws, formula, ref, msg="Pilih dari senarai dropdown."):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error, dv.errorTitle = msg, "Tiada dalam senarai"
    ws.add_data_validation(dv)
    dv.add(ref)


wb = Workbook()

# =====================================================================
# TETAPAN
# =====================================================================
st = wb.active
st.title = "TETAPAN"
title(st, "Tetapan", "Isi sekali. Kalau harga/kos berubah, TAMBAH baris kadar baru dengan tarikh 'Berkuat Dari' — jangan padam yang lama.")
for r, (lab, v, fm) in enumerate([("Tahun laporan", TAHUN, "0"), ("Cukai FB Ads", 0.08, PCT),
                                  ("Cukai Shipping", 0.06, PCT)], start=3):
    cell(st, r, 1, lab, font=f_bold)
    cell(st, r, 2, v, fm, f_in, fill_in)
st["B4"].comment = Comment("Ikut fail asal: TAX FB ADS = FB ADS x 8%", "Sistem")
st["B5"].comment = Comment("Ikut fail asal: TAX SHIPPING = SHIPPING x 6%", "Sistem")
YR, TAX_AD, TAX_SHIP = "TETAPAN!$B$3", "TETAPAN!$B$4", "TETAPAN!$B$5"

# Produk
PH = 8
P1, PN = PH + 1, PH + N_PROD
cell(st, PH - 1, 1, "SENARAI PRODUK", font=f_sec)
ph = ["Produk", "Nama Pakej 1", "Nama Pakej 2", "Nama Pakej 3", "Pcs / Pakej 1", "Pcs / Pakej 2",
      "Pcs / Pakej 3", "Kos Halia Patch (RM/pcs)", "Kos Kayu (RM/pcs)", "Nilai Return Pakej 1 (RM)",
      "Nilai Return Pakej 2 (RM)", "Nilai Return Pakej 3 (RM)"]
for j, h in enumerate(ph, 1):
    head(st, PH, j, h)
prod_rows = [
    ("KRIM EKZEMA", "TRIAL", "JIMAT", "SEMBUH", 1, 3, 5, 2, 5, 39, 69, 89),
    ("KAPSUL RESDUNG", "TRIAL", "JIMAT", "SEMBUH / KELUARGA", 1, 3, 5, 2, 5, 39, 69, 89),
]
for i in range(N_PROD):
    vals = prod_rows[i] if i < len(prod_rows) else (None,) * len(ph)
    for j, v in enumerate(vals, 1):
        cell(st, P1 + i, j, v, RM if j >= 8 else None, f_in, fill_in)
T_PROD = f"TETAPAN!$A${P1}:$A${PN}"

# Saluran
cell(st, PH - 1, 14, "SENARAI SALURAN", font=f_sec)
head(st, PH, 14, "Saluran")
channels = ["WHATSAPP", "ECOMM", "ECOMM SGD", "REPEAT"]
for i in range(N_CH):
    cell(st, P1 + i, 14, channels[i] if i < len(channels) else None, None, f_in, fill_in)
C1 = P1
T_CH = f"TETAPAN!$N${P1}:$N${P1 + N_CH - 1}"

# Kadar
KH = PH + N_PROD + 6          # header row
K1, KN = KH + 1, KH + N_RATE
cell(st, KH - 2, 1, "KADAR HARGA & KOS (ikut produk + saluran)", font=f_sec)
groups = [(1, 3, "KUNCI"), (4, 6, "HARGA ONLINE TRANSFER (RM/pakej)"), (7, 9, "HARGA COD (RM/pakej)"),
          (10, 12, "SHIPPING ONLINE (RM/pakej)"), (13, 15, "SHIPPING COD (RM/pakej)"), (16, 16, "KOS"),
          (17, 19, "FREE HALIA PATCH (pcs/pakej) - OT"), (20, 22, "FREE HALIA PATCH (pcs/pakej) - COD"),
          (23, 25, "FREE KAYU (pcs/pakej) - OT"), (26, 28, "FREE KAYU (pcs/pakej) - COD"), (29, 29, "JUAL")]
for a, b, lab in groups:
    if b > a:
        st.merge_cells(start_row=KH - 1, start_column=a, end_row=KH - 1, end_column=b)
    head(st, KH - 1, a, lab, fill_auto_h)
kh = ["Produk", "Saluran", "Berkuat Dari"] + PAKEJ * 9 + ["Kos Produk (RM/pcs)"]
kh = (["Produk", "Saluran", "Berkuat Dari"] + PAKEJ * 4 + ["Kos Produk (RM/pcs)"] + PAKEJ * 4
      + ["Harga Halia Patch Dijual (RM/pcs)"])
for j, h in enumerate(kh, 1):
    head(st, KH, j, h)
st.row_dimensions[KH - 1].height = 30
st.row_dimensions[KH].height = 30
D = dt.date(TAHUN, 1, 1)
z3 = (0, 0, 0)
# produk, saluran, tarikh, OT, COD, posOT, posCOD, kos, patchOT, patchCOD, kayuOT, kayuCOD, patch jual
rates = [
    ("KRIM EKZEMA", "WHATSAPP", D, (35, 59, 79), (39, 59, 79), (5, 5, 5), (8, 8, 8), 4.5, (0, 0, 2), (0, 0, 2), z3, z3, 0),
    ("KRIM EKZEMA", "WHATSAPP", dt.date(TAHUN, 2, 1), (39, 59, 79), (39, 59, 79), (5, 5, 5), (8, 8, 8), 4.5, (0, 0, 2), (0, 0, 2), z3, z3, 0),
    ("KRIM EKZEMA", "ECOMM", D, (39, 59, 79), (39, 59, 79), (5, 5, 5), (8, 8, 8), 4.5, z3, z3, z3, z3, 15.6),
    ("KRIM EKZEMA", "ECOMM SGD", D, (75, 140, 180), (75, 140, 180), (5, 5, 5), (20, 20, 20), 4.5, z3, z3, z3, z3, 15.6),
    ("KRIM EKZEMA", "REPEAT", D, (39, 59, 79), (39, 59, 79), (5, 5, 5), (8, 8, 8), 5, (0, 0, 1), (0, 0, 1), z3, z3, 0),
    ("KAPSUL RESDUNG", "WHATSAPP", D, (35, 65, 85), (35, 65, 85), (5, 5, 5), (5.85, 6.75, 7.35), 3.7, (0, 1, 2), (0, 1, 2), z3, (0, 0, 1), 0),
    ("KAPSUL RESDUNG", "ECOMM", D, (39, 69, 95), (39, 69, 95), (5, 5, 5), (8, 8, 8), 5, z3, z3, z3, z3, 15.6),
    ("KAPSUL RESDUNG", "REPEAT", D, (35, 65, 85), (39, 69, 89), (5, 5, 5), (8, 8, 8), 5, (0, 1, 2), (0, 1, 2), z3, z3, 0),
]
for i in range(N_RATE):
    r = K1 + i
    if i < len(rates):
        p, s, d, ot, cod, pot, pcod, kos, pa, pb, ka, kb, hj = rates[i]
        vals = [p, s, d, *ot, *cod, *pot, *pcod, kos, *pa, *pb, *ka, *kb, hj]
    else:
        vals = [None] * 29
    for j, v in enumerate(vals, 1):
        fm = DATE if j == 3 else (NUM if 17 <= j <= 28 else (RM if j >= 4 else None))
        cell(st, r, j, v, fm, f_in, fill_in)
st.cell(row=K1 + 1, column=3).comment = Comment(
    "Contoh perubahan harga: Krim Ekzema WhatsApp naik harga Online Transfer Trial 35 -> 39 mula 1 Feb 2026 "
    "(diambil dari fail asal). Order sebelum 1 Feb guna baris atas.", "Sistem")
dv_list(st, f"={T_PROD}", f"A{K1}:A{KN}")
dv_list(st, f"={T_CH}", f"B{K1}:B{KN}")
widths = {1: 18, 2: 14, 3: 14, 14: 14}
for c in range(1, 30):
    st.column_dimensions[L(c)].width = widths.get(c, 10)
st.freeze_panes = f"D{K1}"


def K(col):  # kadar column range
    c = L(col)
    return f"TETAPAN!${c}${K1}:${c}${KN}"


def PR(col):  # produk column range
    c = L(col)
    return f"TETAPAN!${c}${P1}:${c}${PN}"


# =====================================================================
# INPUT HARIAN
# =====================================================================
ih = wb.create_sheet("INPUT HARIAN")
title(ih, f"Input Harian {TAHUN}",
      "1 baris = 1 hari x 1 produk x 1 saluran. Isi kolum KUNING sahaja; kolum biru dikira automatik dari TETAPAN.")
IH = 5
I1, IN = IH + 1, IH + N_INPUT
cols = [  # (label, width, input?)
    ("DATE", 11, 1), ("DAY", 10, 0), ("PRODUK", 16, 1), ("SALURAN", 12, 1),
    ("OT\nTRIAL", 6, 1), ("OT\nJIMAT", 6, 1), ("OT\nSEMBUH /\nKELUARGA", 8, 1),
    ("COD\nTRIAL", 6, 1), ("COD\nJIMAT", 6, 1), ("COD\nSEMBUH /\nKELUARGA", 8, 1),
    ("HALIA PATCH DIJUAL (PCS)", 8, 1), ("FB ADS (RM)", 9, 1), ("LINK CLICK", 7, 1),
    ("FREE POSTAGE (RM)", 8, 1), ("SPECIAL DISCOUNT / ORDER (RM)", 9, 1), ("OPEN AI COST (RM)", 8, 1),
    ("SABAH / SARAWAK (RM)", 8, 1), ("CATATAN", 16, 1),
    ("TOTAL PURCHASE", 8, 0), ("QUANTITY (PCS)", 8, 0), ("FREE HALIA PATCH", 7, 0), ("FREE KAYU", 6, 0),
    ("SALES ONLINE (RM)", 10, 0), ("TAX FB ADS (RM)", 8, 0), ("SHIPPING ONLINE (RM)", 9, 0),
    ("SHIPPING COD (RM)", 9, 0), ("TAX SHIPPING ONLINE (RM)", 8, 0), ("TAX SHIPPING COD (RM)", 8, 0),
    ("COST PRODUCT (RM)", 9, 0), ("FREE GIFT COST (RM)", 8, 0), ("NET PROFIT (RM)", 10, 0),
    ("ROAS", 7, 0), ("ROAS INCLUDING TAX", 7, 0), ("COST PER PURCHASE", 8, 0), ("CONVERSION RATE", 8, 0),
    ("BULAN", 6, 0), ("BARIS KADAR", 6, 0), ("BARIS PRODUK", 6, 0),
]
for j, (lab, w, inp) in enumerate(cols, 1):
    head(ih, IH, j, lab, fill_in_h if inp else fill_auto_h)
    ih.column_dimensions[L(j)].width = w
ih.row_dimensions[IH].height = 52
ih.merge_cells(start_row=IH - 1, start_column=5, end_row=IH - 1, end_column=7)
ih.merge_cells(start_row=IH - 1, start_column=8, end_row=IH - 1, end_column=10)
head(ih, IH - 1, 5, "ONLINE TRANSFER (bil. pakej)")
head(ih, IH - 1, 8, "COD (bil. pakej)")
ih["A3"] = "Header biru tua = ISI  |  Header hijau = AUTO. 'SEMAK KADAR' = tiada kadar dalam TETAPAN untuk produk/saluran/tarikh itu."
ih["A3"].font = f_sub
c_ = {lab.split(" (")[0].replace("\n", " "): L(j) for j, (lab, w, i) in enumerate(cols, 1)}

kp, ks, kd = K(1), K(2), K(3)
for r in range(I1, IN + 1):
    ak = f"$AK{r}"
    al = f"$AL{r}"

    def kr(col):
        return f"INDEX({K(col)},{ak})"

    def pr(col):
        return f"INDEX({PR(col)},{al})"

    ok = f"OR({ak}=0,{al}=0)"
    pak = [f"E{r}", f"F{r}", f"G{r}", f"H{r}", f"I{r}", f"J{r}"]

    def dot(cols6):  # sum of package counts x rates
        return "+".join(f"N({p})*{kr(c)}" for p, c in zip(pak, cols6))

    unit = [pr(5), pr(6), pr(7)] * 2
    f = {
        2: f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r}),"SUNDAY","MONDAY","TUESDAY","WEDNESDAY","THURSDAY","FRIDAY","SATURDAY"))',
        19: f'=IF({ok},"",SUM(E{r}:J{r}))',
        20: f'=IF({ok},"",' + "+".join(f"N({p})*{u}" for p, u in zip(pak, unit)) + ")",
        21: f'=IF({ok},"",{dot([17, 18, 19, 20, 21, 22])})',
        22: f'=IF({ok},"",{dot([23, 24, 25, 26, 27, 28])})',
        23: f'=IF({ok},IF(AND(A{r}<>"",C{r}<>"",D{r}<>""),"SEMAK KADAR",""),{dot([4, 5, 6, 7, 8, 9])}+N(K{r})*{kr(29)}-N(N{r}))',
        24: f'=IF({ok},"",N(L{r})*{TAX_AD})',
        25: f'=IF({ok},"",{"+".join(f"N({p})*{kr(c)}" for p, c in zip(pak[:3], [10, 11, 12]))})',
        26: f'=IF({ok},"",{"+".join(f"N({p})*{kr(c)}" for p, c in zip(pak[3:], [13, 14, 15]))})',
        27: f'=IF({ok},"",Y{r}*{TAX_SHIP})',
        28: f'=IF({ok},"",Z{r}*{TAX_SHIP})',
        29: f'=IF({ok},"",T{r}*{kr(16)})',
        30: f'=IF({ok},"",U{r}*{pr(8)}+V{r}*{pr(9)})',
        31: f'=IF({ok},"",W{r}-N(L{r})-X{r}-Y{r}-Z{r}-AA{r}-AB{r}-AC{r}-AD{r}-N(O{r})-N(P{r})-N(Q{r}))',
        32: f'=IF(OR({ok},N(L{r})=0),"",W{r}/L{r})',
        33: f'=IF(OR({ok},N(L{r})=0),"",W{r}/(L{r}+X{r}))',
        34: f'=IF(OR({ok},N(L{r})=0,N(S{r})=0),"",L{r}/S{r})',
        35: f'=IF(OR({ok},N(M{r})=0),"",S{r}/M{r})',
        36: f'=IF(A{r}="","",IF(YEAR(A{r})={YR},MONTH(A{r}),0))',
        37: f'=IF(OR(A{r}="",C{r}="",D{r}=""),0,SUMPRODUCT(MAX(({kp}=C{r})*({ks}=D{r})*({kd}<=A{r})*(ROW({kp})-{K1 - 1}))))',
        38: f'=IF(C{r}="",0,IFERROR(MATCH(C{r},{T_PROD},0),0))',
    }
    for j, (lab, w, inp) in enumerate(cols, 1):
        x = ih.cell(row=r, column=j, value=f.get(j))
        fm = None
        if j == 1:
            fm = DATE
        elif j in (5, 6, 7, 8, 9, 10, 11, 13, 19, 20, 21, 22):
            fm = NUM
        elif j in (12, 14, 15, 16, 17) or 23 <= j <= 31 or j == 34:
            fm = RM
        elif j in (32, 33):
            fm = X2
        elif j == 35:
            fm = PCT
        x.font = f_in if inp else (f_bold if j == 31 else f_body)
        x.fill = fill_in if inp else fill_auto
        x.border = box
        if fm:
            x.number_format = fm

dv_list(ih, f"={T_PROD}", f"C{I1}:C{IN}")
dv_list(ih, f"={T_CH}", f"D{I1}:D{IN}")
dv = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2100,12,31)", allow_blank=True)
dv.error = "Masukkan tarikh, contoh 15/01/2026"
ih.add_data_validation(dv)
dv.add(f"A{I1}:A{IN}")
dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
dv.error = "Masukkan nombor 0 atau lebih"
ih.add_data_validation(dv)
dv.add(f"E{I1}:Q{IN}")
ih.conditional_formatting.add(f"W{I1}:W{IN}", CellIsRule(operator="equal", formula=['"SEMAK KADAR"'], font=red,
                                                         fill=PatternFill("solid", fgColor="F8CBAD")))
ih.conditional_formatting.add(f"AE{I1}:AE{IN}", CellIsRule(operator="lessThan", formula=["0"], font=red))
ih.conditional_formatting.add(f"A{I1}:A{IN}", FormulaRule(formula=[f'AND($A{I1}<>"",YEAR($A{I1})<>{YR})'],
                                                          fill=PatternFill("solid", fgColor="F8CBAD")))
ih.freeze_panes = f"E{I1}"
ih.auto_filter.ref = f"A{IH}:AL{IN}"
for c in ("AJ", "AK", "AL"):
    ih.column_dimensions[c].hidden = True


def IR(col):
    return f"'INPUT HARIAN'!${col}${I1}:${col}${IN}"


R_DATE, R_PROD, R_CH, R_BLN = IR("A"), IR("C"), IR("D"), IR("AJ")
R_SALES, R_ADS, R_TAXAD, R_NP = IR("W"), IR("L"), IR("X"), IR("AE")
R_PUR, R_PCS, R_CLICK = IR("S"), IR("T"), IR("M")

# =====================================================================
# RETURN
# =====================================================================
rt = wb.create_sheet("RETURN")
title(rt, f"Return {TAHUN}", "Rekod setiap return. Nilai return ditolak dari NET PROFIT bulan berkenaan.")
RH = 4
R1_, RN_ = RH + 1, RH + N_RETURN
for j, (lab, w, inp) in enumerate([("DATE", 11, 1), ("PRODUK", 18, 1), ("PAKEJ", 11, 1), ("BILANGAN", 9, 1),
                                   ("SPECIAL ORDER (RM)", 11, 1), ("CATATAN", 24, 1),
                                   ("TOTAL RETURN (RM)", 12, 0), ("BULAN", 7, 0)], 1):
    head(rt, RH, j, lab, fill_in_h if inp else fill_auto_h)
    rt.column_dimensions[L(j)].width = w
rt.row_dimensions[RH].height = 30
for r in range(R1_, RN_ + 1):
    for j in range(1, 7):
        cell(rt, r, j, None, DATE if j == 1 else (RM if j == 5 else (NUM if j == 4 else None)), f_in, fill_in)
    cell(rt, r, 7, f'=IF(B{r}="","",IFERROR(N(D{r})*INDEX(TETAPAN!$J${P1}:$L${PN},MATCH(B{r},{T_PROD},0),'
                    f'IF(C{r}="TRIAL",1,IF(C{r}="JIMAT",2,3))),0)+N(E{r}))', RM, fill=fill_auto)
    cell(rt, r, 8, f'=IF(A{r}="","",IF(YEAR(A{r})={YR},MONTH(A{r}),0))', fill=fill_auto)
rt.column_dimensions["H"].hidden = True
dv_list(rt, f"={T_PROD}", f"B{R1_}:B{RN_}")
dv_list(rt, '"TRIAL,JIMAT,SEMBUH"', f"C{R1_}:C{RN_}")
rt["J4"], rt["J5"] = "Nota", "Pakej SEMBUH = pakej ke-3 (SEMBUH / KELUARGA). Nilai seunit ikut TETAPAN > Nilai Return."
rt["J4"].font, rt["J5"].font = f_bold, f_sub
rt.freeze_panes = f"A{R1_}"
RT_P, RT_V, RT_B, RT_D = (f"RETURN!${c}${R1_}:${c}${RN_}" for c in "BGHA")

# =====================================================================
# RINGKASAN BULANAN (ganti TOTAL 2026)
# =====================================================================
rb = wb.create_sheet("RINGKASAN BULANAN")
title(rb, f"Ringkasan Bulanan {TAHUN}", "Automatik sepenuhnya (ganti sheet TOTAL 2026). Jangan taip di sini.")
cell(rb, 3, 2, "TOTAL CUMULATIVE SALES", font=f_bold, fill=fill_tot)
cell(rb, 3, 5, "TOTAL CUMULATIVE PROFIT", font=f_bold, fill=fill_tot)
row = 6
sec_rows = []          # (title_row, first_month_row) per product
for p in range(N_PROD):
    tr = row
    x = rb.cell(row=tr, column=1, value=f'=IF(TETAPAN!A{P1 + p}="","(slot produk kosong)",TETAPAN!A{P1 + p})')
    x.font = f_sec
    hr = tr + 1
    heads = ["MONTH"]
    heads += [f'=IF(TETAPAN!$N${C1 + k}="","-","SALES "&TETAPAN!$N${C1 + k})' for k in range(N_CH)] + ["TOTAL SALES"]
    heads += [f'=IF(TETAPAN!$N${C1 + k}="","-","ADS "&TETAPAN!$N${C1 + k})' for k in range(N_CH)] + ["TOTAL ADS"]
    heads += [f'=IF(TETAPAN!$N${C1 + k}="","-","PROFIT "&TETAPAN!$N${C1 + k})' for k in range(N_CH)]
    heads += ["TOTAL RETURN", "NET PROFIT", "PROFIT MARGIN", "TOTAL PURCHASE", "QUANTITY (PCS)", "ROAS"]
    for j, h in enumerate(heads, 1):
        head(rb, hr, j, h)
    rb.row_dimensions[hr].height = 34
    pref = f"$A${tr}"
    for m in range(12):
        r = hr + 1 + m
        cell(rb, r, 1, BULAN[m], font=f_bold)
        crit = f"{R_PROD},{pref},{R_BLN},{m + 1}"
        for k in range(N_CH):
            ch = f"TETAPAN!$N${C1 + k}"
            cell(rb, r, 2 + k, f'=IF({ch}="",0,SUMIFS({R_SALES},{crit},{R_CH},{ch}))', RM)
            cell(rb, r, 8 + k, f'=IF({ch}="",0,SUMIFS({R_ADS},{crit},{R_CH},{ch}))', RM)
            cell(rb, r, 14 + k, f'=IF({ch}="",0,SUMIFS({R_NP},{crit},{R_CH},{ch}))', RM)
        cell(rb, r, 7, f"=SUM(B{r}:F{r})", RM, f_bold)
        cell(rb, r, 13, f"=SUM(H{r}:L{r})", RM, f_bold)
        cell(rb, r, 19, f"=SUMIFS({RT_V},{RT_P},{pref},{RT_B},{m + 1})", RM)
        cell(rb, r, 20, f"=SUM(N{r}:R{r})-S{r}", RM, f_bold)
        cell(rb, r, 21, f"=IF(G{r}=0,0,T{r}/G{r})", PCT)
        cell(rb, r, 22, f"=SUMIFS({R_PUR},{crit})", NUM)
        cell(rb, r, 23, f"=SUMIFS({R_PCS},{crit})", NUM)
        cell(rb, r, 24, f"=IF(M{r}=0,0,G{r}/M{r})", X2)
    tot = hr + 13
    cell(rb, tot, 1, "TOTAL", font=f_bold, fill=fill_tot)
    for j in range(2, 25):
        cl = L(j)
        if j == 21:
            v = f"=IF(G{tot}=0,0,T{tot}/G{tot})"
        elif j == 24:
            v = f"=IF(M{tot}=0,0,G{tot}/M{tot})"
        else:
            v = f"=SUM({cl}{hr + 1}:{cl}{hr + 12})"
        cell(rb, tot, j, v, rb.cell(row=hr + 1, column=j).number_format, f_bold, fill_tot)
    sec_rows.append((tr, hr + 1, tot))
    row = tot + 3

# Keseluruhan
ktr = row
rb.cell(row=ktr, column=1, value="KESELURUHAN (SEMUA PRODUK)").font = f_sec
khr = ktr + 1
kheads = ["MONTH", "TOTAL SALES", "TOTAL ADS", "TAX ADS", "TOTAL RETURN", "NET PROFIT", "PROFIT MARGIN",
          "TOTAL PURCHASE", "QUANTITY (PCS)", "ROAS", "CUMULATIVE SALES", "CUMULATIVE PROFIT", "SALES GROWTH"]
for j, h in enumerate(kheads, 1):
    head(rb, khr, j, h)
rb.row_dimensions[khr].height = 30
for m in range(12):
    r = khr + 1 + m
    cell(rb, r, 1, BULAN[m], font=f_bold)

    def acr(col):
        return "+".join(f"{col}{s[1] + m}" for s in sec_rows)

    cell(rb, r, 2, f"={acr('G')}", RM, f_bold)
    cell(rb, r, 3, f"={acr('M')}", RM)
    cell(rb, r, 4, f"=SUMIFS({R_TAXAD},{R_BLN},{m + 1})", RM)
    cell(rb, r, 5, f"={acr('S')}", RM)
    cell(rb, r, 6, f"={acr('T')}", RM, f_bold)
    cell(rb, r, 7, f"=IF(B{r}=0,0,F{r}/B{r})", PCT)
    cell(rb, r, 8, f"={acr('V')}", NUM)
    cell(rb, r, 9, f"={acr('W')}", NUM)
    cell(rb, r, 10, f"=IF(C{r}=0,0,B{r}/C{r})", X2)
    cell(rb, r, 11, f"=SUM($B${khr + 1}:B{r})", RM)
    cell(rb, r, 12, f"=SUM($F${khr + 1}:F{r})", RM)
    cell(rb, r, 13, "=0" if m == 0 else f"=IF(B{r - 1}=0,0,B{r}/B{r - 1}-1)", PCT)
ktot = khr + 13
cell(rb, ktot, 1, "TOTAL", font=f_bold, fill=fill_tot)
for j in range(2, 14):
    cl = L(j)
    v = {7: f"=IF(B{ktot}=0,0,F{ktot}/B{ktot})", 10: f"=IF(C{ktot}=0,0,B{ktot}/C{ktot})",
         11: "", 12: "", 13: ""}.get(j, f"=SUM({cl}{khr + 1}:{cl}{khr + 12})")
    cell(rb, ktot, j, v, rb.cell(row=khr + 1, column=j).number_format, f_bold, fill_tot)
rb.merge_cells("B4:C4")
rb.merge_cells("E4:F4")
cell(rb, 4, 2, f"=B{ktot}", RM, Font(name=F, size=14, bold=True, color=NAVY))
cell(rb, 4, 5, f"=F{ktot}", RM, Font(name=F, size=14, bold=True, color=GREEN))
rb.conditional_formatting.add(f"A6:X{ktot}", CellIsRule(operator="lessThan", formula=["0"], font=red))
rb.cell(row=ktot + 1, column=1,
        value="TOTAL ADS = FB Ads sebelum cukai (cukai di kolum TAX ADS). NET PROFIT sudah tolak semua kos, cukai & return.").font = f_sub
rb.column_dimensions["A"].width = 14
for j in range(2, 25):
    rb.column_dimensions[L(j)].width = 11.5
rb.freeze_panes = "B6"

ch1 = BarChart()
ch1.title = "Total Sales vs Net Profit (Bulanan)"
ch1.height, ch1.width, ch1.y_axis.title = 8, 22, "RM"
ch1.add_data(Reference(rb, min_col=2, min_row=khr, max_row=khr + 12), titles_from_data=True)
ch1.add_data(Reference(rb, min_col=6, min_row=khr, max_row=khr + 12), titles_from_data=True)
ch1.set_categories(Reference(rb, min_col=1, min_row=khr + 1, max_row=khr + 12))
rb.add_chart(ch1, f"A{ktot + 3}")
ch2 = LineChart()
ch2.title = "Total Ads (Bulanan)"
ch2.height, ch2.width = 8, 14
ch2.add_data(Reference(rb, min_col=3, min_row=khr, max_row=khr + 12), titles_from_data=True)
ch2.set_categories(Reference(rb, min_col=1, min_row=khr + 1, max_row=khr + 12))
rb.add_chart(ch2, f"M{ktot + 3}")

# =====================================================================
# PROFIT HARIAN (ganti TOTALPROFIT/DAY)
# =====================================================================
dh = wb.create_sheet("PROFIT HARIAN")
title(dh, f"Profit Harian {TAHUN}", "Total sales, ads & net profit setiap hari ikut produk (automatik).")
DH = 5
hd = ["DATE", "DAY", "BULAN"]
for p in range(N_PROD):
    for lab in ("SALES", "ADS", "PROFIT"):
        hd.append(f'=IF(TETAPAN!$A${P1 + p}="","-",TETAPAN!$A${P1 + p}&" {lab}")')
hd += ["RETURN (RM)", "TOTAL SALES", "TOTAL ADS", "TOTAL PROFIT / DAY"]
for j, h in enumerate(hd, 1):
    head(dh, DH, j, h, fill_auto_h if j > 3 + 3 * N_PROD else fill_in_h)
dh.row_dimensions[DH].height = 42
ndays = 366 if calendar.isleap(TAHUN) else 365
for i in range(ndays):
    r = DH + 1 + i
    cell(dh, r, 1, f"=DATE({YR},1,1)+{i}", DATE, f_bold)
    cell(dh, r, 2, f'=CHOOSE(WEEKDAY(A{r}),"SUNDAY","MONDAY","TUESDAY","WEDNESDAY","THURSDAY","FRIDAY","SATURDAY")')
    cell(dh, r, 3, f"=MONTH(A{r})")
    for p in range(N_PROD):
        pn = f"TETAPAN!$A${P1 + p}"
        base = 4 + 3 * p
        for k, rng_ in enumerate((R_SALES, R_ADS, R_NP)):
            cell(dh, r, base + k, f'=IF({pn}="",0,SUMIFS({rng_},{R_DATE},$A{r},{R_PROD},{pn}))', RM)
    c0 = 4 + 3 * N_PROD
    cell(dh, r, c0, f"=SUMIFS({RT_V},{RT_D},$A{r})", RM)
    cell(dh, r, c0 + 1, "=" + "+".join(f"{L(4 + 3 * p)}{r}" for p in range(N_PROD)), RM, f_bold)
    cell(dh, r, c0 + 2, "=" + "+".join(f"{L(5 + 3 * p)}{r}" for p in range(N_PROD)), RM)
    cell(dh, r, c0 + 3, "=" + "+".join(f"{L(6 + 3 * p)}{r}" for p in range(N_PROD)) + f"-{L(c0)}{r}", RM, f_bold)
dtot = DH + ndays + 1
cell(dh, dtot, 1, "TOTAL", font=f_bold, fill=fill_tot)
for j in range(4, c0 + 4):
    cell(dh, dtot, j, f"=SUM({L(j)}{DH + 1}:{L(j)}{dtot - 1})", RM, f_bold, fill_tot)
dh["A3"] = "Guna butang penapis (filter) di kolum BULAN untuk lihat satu bulan sahaja."
dh["A3"].font = f_sub
dh.conditional_formatting.add(f"D{DH + 1}:{L(c0 + 3)}{dtot}", CellIsRule(operator="lessThan", formula=["0"], font=red))
dh.auto_filter.ref = f"A{DH}:{L(c0 + 3)}{dtot - 1}"
dh.freeze_panes = f"D{DH + 1}"
dh.column_dimensions["A"].width = 11
dh.column_dimensions["B"].width = 11
dh.column_dimensions["C"].width = 7
for j in range(4, c0 + 4):
    dh.column_dimensions[L(j)].width = 11.5

# =====================================================================
# DASHBOARD
# =====================================================================
db = wb.create_sheet("DASHBOARD")
title(db, f"Dashboard P&L {TAHUN}", "Pilih bulan di C4. Semua angka automatik.")
cell(db, 4, 2, "PILIH BULAN:", font=f_bold)
cell(db, 4, 3, BULAN[0], font=Font(name=F, size=12, bold=True, color="0000FF"), fill=fill_in, align=center)
RBn = "'RINGKASAN BULANAN'!"
dv_list(db, '"' + ",".join(BULAN) + '"', "C4")
db["D4"] = f"=MATCH(C4,{RBn}$A${khr + 1}:$A${khr + 12},0)"
db["D4"].font = Font(name=F, size=8, color="FFFFFF")
kpi = [("TOTAL SALES", "B", RM), ("NET PROFIT", "F", RM), ("PROFIT MARGIN", "G", PCT), ("TOTAL ADS", "C", RM),
       ("ROAS", "J", X2), ("TOTAL RETURN", "E", RM), ("TOTAL PURCHASE", "H", NUM), ("QUANTITY (PCS)", "I", NUM)]
for j, h in enumerate(["", "PETUNJUK", "BULAN DIPILIH", "BULAN SEBELUM", "BEZA %", "SETAHUN (YTD)"], 1):
    head(db, 6, j, h)
for i, (lab, col, fm) in enumerate(kpi):
    r = 7 + i
    rng_ = f"{RBn}${col}${khr + 1}:${col}${khr + 12}"
    cell(db, r, 2, lab, font=f_bold)
    cell(db, r, 3, f"=INDEX({rng_},$D$4)", fm, Font(name=F, size=11, bold=True, color=NAVY))
    cell(db, r, 4, f"=IF($D$4=1,0,INDEX({rng_},$D$4-1))", fm)
    cell(db, r, 5, f"=IF(D{r}=0,0,C{r}/D{r}-1)", PCT)
    cell(db, r, 6, f"={RBn}{col}{ktot}", fm, f_bold)
db.conditional_formatting.add("E7:E14", CellIsRule(operator="lessThan", formula=["0"], font=red))
db.conditional_formatting.add("E7:E14", CellIsRule(operator="greaterThan", formula=["0"],
                                                   font=Font(name=F, size=10, bold=True, color=GREEN)))
# produk x saluran untuk bulan dipilih
TR0 = 17
db.cell(row=TR0 - 1, column=2, value="PRESTASI IKUT PRODUK & SALURAN (BULAN DIPILIH)").font = f_sec
for j, h in enumerate(["", "PRODUK", "SALURAN", "SALES (RM)", "FB ADS (RM)", "NET PROFIT (RM)", "ROAS",
                       "PURCHASE", "PCS", "LINK CLICK", "CONVERSION", "COST / PURCHASE"], 1):
    head(db, TR0, j, h)
db.row_dimensions[TR0].height = 30
r = TR0 + 1
for p in range(N_PROD):
    for k in range(N_CH):
        pn, cn = f"TETAPAN!$A${P1 + p}", f"TETAPAN!$N${C1 + k}"
        cell(db, r, 2, f'=IF(OR({pn}="",{cn}=""),"",{pn})')
        cell(db, r, 3, f'=IF(B{r}="","",{cn})')
        crit = f"{R_PROD},B{r},{R_CH},C{r},{R_BLN},$D$4"
        for j, rng_, fm in ((4, R_SALES, RM), (5, R_ADS, RM), (6, R_NP, RM), (8, R_PUR, NUM), (9, R_PCS, NUM),
                            (10, R_CLICK, NUM)):
            cell(db, r, j, f'=IF(B{r}="","",SUMIFS({rng_},{crit}))', fm)
        cell(db, r, 7, f'=IF(N(E{r})=0,"",D{r}/E{r})', X2)
        cell(db, r, 11, f'=IF(N(J{r})=0,"",H{r}/J{r})', PCT)
        cell(db, r, 12, f'=IF(N(H{r})=0,"",E{r}/H{r})', RM)
        r += 1
cell(db, r, 2, "TOTAL", font=f_bold, fill=fill_tot)
cell(db, r, 3, None, fill=fill_tot)
for j in (4, 5, 6, 8, 9, 10):
    cell(db, r, j, f"=SUM({L(j)}{TR0 + 1}:{L(j)}{r - 1})", RM if j < 8 else NUM, f_bold, fill_tot)
cell(db, r, 7, f'=IF(E{r}=0,"",D{r}/E{r})', X2, f_bold, fill_tot)
cell(db, r, 11, f'=IF(J{r}=0,"",H{r}/J{r})', PCT, f_bold, fill_tot)
cell(db, r, 12, f'=IF(H{r}=0,"",E{r}/H{r})', RM, f_bold, fill_tot)
db.conditional_formatting.add(f"F{TR0 + 1}:F{r}", CellIsRule(operator="lessThan", formula=["0"], font=red))
for c, w in {"A": 3, "B": 18, "C": 15, "D": 14, "E": 14, "F": 15, "G": 9, "H": 10, "I": 9, "J": 10, "K": 11, "L": 12}.items():
    db.column_dimensions[c].width = w
ch3 = BarChart()
ch3.title = "Sales vs Net Profit"
ch3.height, ch3.width = 7.5, 16
ch3.add_data(Reference(rb, min_col=2, min_row=khr, max_row=khr + 12), titles_from_data=True)
ch3.add_data(Reference(rb, min_col=6, min_row=khr, max_row=khr + 12), titles_from_data=True)
ch3.set_categories(Reference(rb, min_col=1, min_row=khr + 1, max_row=khr + 12))
db.add_chart(ch3, "H3")

# =====================================================================
# PANDUAN
# =====================================================================
pg = wb.create_sheet("PANDUAN", 0)
title(pg, f"Sistem P&L Jualan Online {TAHUN}", "Baca sekali sebelum mula.")
lines = [
    ("SETIAP HARI", None),
    ("1", "Buka INPUT HARIAN. Satu baris untuk setiap PRODUK + SALURAN yang ada jualan/iklan hari itu."),
    ("2", "Isi: DATE, PRODUK (dropdown), SALURAN (dropdown), bilangan pakej OT & COD, FB ADS, LINK CLICK."),
    ("3", "Kalau ada: Halia Patch dijual (Ecomm), Free Postage, Special Discount/Order, Open AI Cost, Sabah/Sarawak."),
    ("4", "Semua kolum hijau (sales, shipping, cukai, kos, net profit, ROAS...) keluar sendiri."),
    ("5", "Ada return? Rekod di sheet RETURN (date, produk, pakej, bilangan)."),
    ("", ""),
    ("RESULT AUTOMATIK", None),
    ("DASHBOARD", "Pilih bulan: KPI bulan itu vs bulan lepas + YTD, dan prestasi setiap produk x saluran."),
    ("RINGKASAN BULANAN", "Ganti TOTAL 2026: sales/ads/profit ikut saluran, return, net profit, margin, ROAS, kumulatif."),
    ("PROFIT HARIAN", "Ganti TOTALPROFIT/DAY: sales, ads & profit setiap hari ikut produk."),
    ("", ""),
    ("TUKAR HARGA / KOS", None),
    ("TETAPAN", "Harga pakej, shipping, kos produk, free gift ikut PRODUK + SALURAN. Harga berubah? Tambah baris baru"),
    ("", "dengan tarikh 'Berkuat Dari'. Order lama kekal guna harga lama, order baru guna harga baru."),
    ("", "Produk/saluran baru: tambah di SENARAI PRODUK / SENARAI SALURAN, kemudian tambah kadar untuknya."),
    ("", ""),
    ("FORMULA (sama macam fail manual)", None),
    ("QUANTITY", "Pakej x pcs/pakej (Trial 1, Jimat 3, Sembuh/Keluarga 5)."),
    ("SALES", "Pakej x harga + halia patch dijual x harga patch - free postage."),
    ("TAX", "FB Ads x 8%; Shipping x 6%."),
    ("NET PROFIT", "Sales - FB Ads - tax ads - shipping OT & COD - tax shipping - cost product - free gift cost"),
    ("", "- special discount/order - open AI - Sabah/Sarawak."),
    ("ROAS", "Sales / FB Ads.  ROAS incl. tax = Sales / (FB Ads + tax).  Conversion = Purchase / Link Click."),
    ("", ""),
    ("WARNA", None),
    ("Kuning", "Sel untuk diisi."),
    ("Biru muda", "Formula - jangan taip."),
    ("SEMAK KADAR", "Muncul jika tiada kadar di TETAPAN untuk produk/saluran/tarikh itu."),
]
for i, (a, b) in enumerate(lines, 4):
    ca, cb = pg.cell(row=i, column=1, value=a), pg.cell(row=i, column=2, value=b)
    if b is None:
        ca.font = Font(name=F, size=11, bold=True, color=TEAL)
    else:
        ca.font, cb.font = f_bold, f_body
    if a == "Kuning":
        ca.fill = fill_in
    if a == "Biru muda":
        ca.fill = fill_auto
pg.column_dimensions["A"].width = 22
pg.column_dimensions["B"].width = 110

# =====================================================================
# Pindah data dari fail manual
# =====================================================================
if len(sys.argv) > 2:
    src = load_workbook(sys.argv[2])
    # blok: (produk, saluran, baris_mula, kolum pakej x6, peta kolum lain -> kolum INPUT)
    blocks = [
        ("KRIM EKZEMA", "WHATSAPP", 7, "D E F G H I", {"N": "L", "AC": "M", "L": "N", "V": "O", "W": "P"}),
        ("KAPSUL RESDUNG", "WHATSAPP", 7, "AI AJ AK AL AM AN", {"AS": "L", "BI": "M", "BB": "O", "BC": "P", "BA": "Q"}),
        ("KRIM EKZEMA", "ECOMM", 47, "D E F G H I", {"J": "K", "M": "L", "T": "O"}),
        ("KAPSUL RESDUNG", "ECOMM", 47, "AI AJ AK AL AM AN", {"AO": "K", "AR": "L", "AY": "O"}),
        ("KRIM EKZEMA", "ECOMM SGD", 86, "D E F G H I", {"J": "K", "M": "L", "T": "O"}),
        ("KAPSUL RESDUNG", "REPEAT", 89, "AI AJ AK AL AM AN", {"AQ": "N", "AY": "O"}),
        ("KRIM EKZEMA", "REPEAT", 125, "D E F G H I", {"L": "N", "T": "O"}),
    ]
    order_p = {p[0]: i for i, p in enumerate(prod_rows)}
    order_c = {c: i for i, c in enumerate(channels)}
    recs = []
    for mi, sname in enumerate(["JAN", "FEB", "MAC", "APR", "MEI", "JUN", "JUL", "OGOS", "SEP", "OKT", "NOV", "DIS"]):
        sname = f"{sname} {TAHUN}"
        if sname not in src.sheetnames:
            continue
        ws = src[sname]
        nd = calendar.monthrange(TAHUN, mi + 1)[1]
        for prod, chan, r0, pk, extra in blocks:
            for d in range(nd):
                r = r0 + d
                vals = {}
                for sc, dc in zip(pk.split(), "EFGHIJ"):
                    vals[dc] = ws[f"{sc}{r}"].value
                for sc, dc in extra.items():
                    vals[dc] = ws[f"{sc}{r}"].value
                vals = {k: v for k, v in vals.items() if isinstance(v, (int, float)) and v != 0}
                if vals:
                    recs.append((dt.date(TAHUN, mi + 1, d + 1), order_p[prod], order_c[chan], prod, chan, vals))
    recs.sort(key=lambda x: x[:3])
    for i, (d, _, _, prod, chan, vals) in enumerate(recs):
        r = I1 + i
        ih[f"A{r}"], ih[f"C{r}"], ih[f"D{r}"] = d, prod, chan
        for k, v in vals.items():
            ih[f"{k}{r}"] = v
    print(f"Dipindah: {len(recs)} baris input dari {sys.argv[2]}")

# susunan sheet
order = ["PANDUAN", "DASHBOARD", "INPUT HARIAN", "RETURN", "RINGKASAN BULANAN", "PROFIT HARIAN", "TETAPAN"]
wb._sheets = [wb[n] for n in order]
wb.active = 2
for n, col in [("PANDUAN", NAVY), ("DASHBOARD", NAVY), ("INPUT HARIAN", "F2C200"), ("RETURN", "F2C200"),
               ("TETAPAN", "F2C200"), ("RINGKASAN BULANAN", TEAL), ("PROFIT HARIAN", TEAL)]:
    wb[n].sheet_properties.tabColor = col
wb.save(sys.argv[1])
print(sys.argv[1])
